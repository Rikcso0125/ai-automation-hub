# -*- coding: utf-8 -*-
import os
import json
import re
from pathlib import Path
from typing import Dict, Any

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document

def slugify(text: str) -> str:
    text = text.lower()
    replacements = {
        'á': 'a', 'é': 'e', 'í': 'i', 'ó': 'o', 'ö': 'o', 'ő': 'o',
        'ú': 'u', 'ü': 'u', 'ű': 'u', ' ': '.'
    }
    for k, v in replacements.items():
        text = text.replace(k, v)
    return re.sub(r'[^a-z0-9.]', '', text)

def generate_contract_pdf(data: Dict[str, Any], output_path: str):
    doc = SimpleDocTemplate(output_path, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#0f172a'),
        alignment=1,
        spaceAfter=14
    )
    
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=15,
        textColor=colors.HexColor('#334155'),
        spaceAfter=8
    )

    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1e293b'),
        spaceBefore=10,
        spaceAfter=4
    )

    story = []
    story.append(Paragraph("MUNKASZERZODES ES TITOKTARTASI NYILATKOZAT", title_style))
    story.append(Spacer(1, 10))

    story.append(Paragraph(f"amely letrejott egyreszrol a(z) <b>{data['company_name']}</b> mint Munkaltato,", body_style))
    story.append(Paragraph(f"masreszrol <b>{data['full_name']}</b> mint Munkavallalo kozott az alabbi feltetelekkel:", body_style))
    story.append(Spacer(1, 10))

    table_data = [
        ["Szerzodeses Pont", "Reszletek"],
        ["Munkakori megnevezes", data['position']],
        ["Szervezeti egyseg", data['department']],
        ["Munkaviszony kezdete", data['start_date']],
        ["Brutto alapber", f"{data['gross_salary_huf']:,} Ft / ho".replace(",", " ")],
        ["Probaido idotartama", "3 (harom) honap"],
        ["Kijelolt szakmai mentor", data['mentor_name']],
        ["Hivatalos ceges email", data['company_email']]
    ]

    t = Table(table_data, colWidths=[160, 320])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0284c7')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 9.5),
        ('BOTTOMPADDING', (0,0), (-1,0), 5),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('BACKGROUND', (0,1), (-1,-1), colors.HexColor('#f8fafc')),
        ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
        ('FONTSIZE', (0,1), (-1,-1), 9),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(t)
    story.append(Spacer(1, 14))

    story.append(Paragraph("1. MUNKAKORI FELADATOK ES TITOKTARTAS (NDA)", h2_style))
    story.append(Paragraph(
        "A Munkavallalo kotelezettseget vallal arra, hogy a Munkaltato mukodesevel, ugyfeleivel, "
        "technologiai megoldasaival kapcsolatos valamennyi uzleti es technikai titkot bizalmasan kezel, "
        "azokat harmadik fel szamara nem adja at.",
        body_style
    ))
    story.append(Spacer(1, 10))

    story.append(Paragraph("2. ESZKOZHASZNALAT ES IT BIZTONSAG", h2_style))
    story.append(Paragraph(
        "A Munkaltato biztositja a munkavegzeshez szukseges hardver es szoftver eszkozoket. "
        "A Munkavallalo vallalja az IT biztonsagi szabalyzat es a ketlepcsos hitelesites (2FA) betartasat.",
        body_style
    ))

    doc.build(story)

def generate_contract_docx(data: Dict[str, Any], output_path: str):
    doc = Document()
    doc.add_heading("MUNKASZERZŐDÉS ÉS TITOKTARTÁSI NYILATKOZAT", level=1)
    doc.add_paragraph(f"Munkáltató: {data['company_name']}")
    doc.add_paragraph(f"Munkavállaló: {data['full_name']}")
    doc.add_paragraph(f"Munkakör: {data['position']} ({data['department']})")
    doc.add_paragraph(f"Kezdés dátuma: {data['start_date']}")
    doc.add_paragraph(f"Bruttó bér: {data['gross_salary_huf']:,} Ft/hó".replace(",", " "))
    doc.add_paragraph(f"Hivatalos céges email: {data['company_email']}")
    doc.add_paragraph(f"Szakmai mentor: {data['mentor_name']}")
    doc.save(output_path)

def run(payload: Dict[str, Any], config: Dict[str, Any]) -> Dict[str, Any]:
    company_name = config.get("company_name", "ProfiTech Ipari & Kereskedelmi Kft.")
    domain = config.get("company_email_domain", "profitech.hu")
    workspace_sys = config.get("default_workspace_system", "Google Workspace (Gmail, Drive, Meet)")
    comm_tool = config.get("communication_tool", "Slack Workspace")
    pm_sys = config.get("pm_system", "ClickUp")
    cloud_storage = config.get("cloud_storage", "Google Drive Shared Drive")

    onboarding_type = payload.get("onboarding_type", "employee")
    details = payload.get("details", {})

    full_name = details.get("full_name", "Uj Munkatars")
    position = details.get("position", "Munkatars")
    dept = details.get("department", "Operacio")
    start_date = details.get("start_date", "2026-10-01")
    gross_salary = details.get("gross_salary_huf", 1000000)
    mentor = details.get("mentor_name", "Kovacs Anna")
    mentor_email = details.get("mentor_email", f"mentor@{domain}")
    personal_email = details.get("personal_email", "magan@email.hu")
    hardware_list = details.get("hardware_needed", ["Laptop", "Monitor", "Mobil"])

    # 1. Automatikus Céges Email & Felhasználónév Generálás
    email_user = slugify(full_name)
    company_email = f"{email_user}@{domain}"

    # Adatok a szerződéshez
    contract_data = {
        "company_name": company_name,
        "full_name": full_name,
        "position": position,
        "department": dept,
        "start_date": start_date,
        "gross_salary_huf": gross_salary,
        "mentor_name": mentor,
        "company_email": company_email
    }

    # 2. Szerződés & NDA Generálás (PDF és Word)
    out_dir = Path(__file__).resolve().parent / "output"
    out_dir.mkdir(parents=True, exist_ok=True)
    pdf_filename = f"Munkaszerzodes_{email_user}.pdf"
    docx_filename = f"Munkaszerzodes_{email_user}.docx"
    pdf_path = str(out_dir / pdf_filename)
    docx_path = str(out_dir / docx_filename)

    generate_contract_pdf(contract_data, pdf_path)
    generate_contract_docx(contract_data, docx_path)

    # 3. IT & Fiók Létrehozási Teendők (Provisioning)
    provisioning_steps = [
        {
            "system": workspace_sys,
            "account_created": company_email,
            "role": f"Employee ({dept})",
            "status": "PROVISIONED_SUCCESS"
        },
        {
            "system": comm_tool,
            "channel_invites": ["#altalanos", f"#{dept.lower().replace(' ', '-')}", "#onboarding-2026"],
            "status": "INVITATION_SENT"
        },
        {
            "system": pm_sys,
            "seat_assigned": company_email,
            "project_boards": [dept, "Onboarding Checklist"],
            "status": "SEAT_ACTIVATED"
        },
        {
            "system": cloud_storage,
            "dedicated_folder": f"/munkatarsak/{email_user}/",
            "permissions": "Read & Write",
            "status": "FOLDER_CREATED"
        }
    ]

    # 4. IT Eszközigény Ticket
    it_ticket = {
        "ticket_id": f"IT-REQ-{email_user.upper()[:6]}",
        "beneficiary": full_name,
        "deadline": f"{start_date} elott 1 munkanappal",
        "items": hardware_list,
        "assigned_to": config.get("it_support_email", "it@profitech.hu")
    }

    # 5. Drip Onboarding Ütemezés (1., 3. és 7. nap)
    drip_schedule = [
        {
            "day": "1. Nap (Kezdés napja - 08:30)",
            "action": "Üdvözlő Email & Belépési Kulcsok",
            "recipient": personal_email,
            "subject": f"Üdvözlünk a(z) {company_name} csapatában, {full_name}!",
            "summary": f"Belépési adatok a(z) {company_email} fiókhoz, jelszó-beállítási link, mentor ({mentor}) elérhetősége és az első nap menetrendje."
        },
        {
            "day": "3. Nap (Eszközök & Belső Rendszerek - 09:00)",
            "action": "Belső Tudásbázis & Első Feladat",
            "recipient": company_email,
            "subject": "3. Napi Onboarding Mérföldkő: Belső Rendszerek & Képzés",
            "summary": f"Hozzáférés a céges kézikönyvhöz a(z) {cloud_storage} felületen, valamint az első éles feladat a(z) {pm_sys} táblán."
        },
        {
            "day": "7. Nap (Visszajelzés & Vezetői 1-on-1 - 14:00)",
            "action": "1 Hetes Visszajelző Kérdőív & Meeting",
            "recipient": company_email,
            "subject": "Hogy telt az első heted? 1 hetes visszajelzés & 1-on-1 egyeztetés",
            "summary": f"Rövid 5 perces elégedettségi kérdőív kitöltése és 30 perces naptárfoglalás a szakmai mentorral ({mentor})."
        }
    ]

    return {
        "status": "success",
        "onboarding_type": onboarding_type,
        "employee_summary": {
            "name": full_name,
            "position": position,
            "department": dept,
            "start_date": start_date,
            "company_email": company_email,
            "mentor": mentor
        },
        "generated_documents": {
            "pdf_contract": pdf_path,
            "docx_contract": docx_path,
            "pdf_filename": pdf_filename,
            "docx_filename": docx_filename
        },
        "provisioning": {
            "systems_count": len(provisioning_steps),
            "steps": provisioning_steps,
            "it_hardware_ticket": it_ticket
        },
        "drip_schedule": {
            "total_phases": len(drip_schedule),
            "phases": drip_schedule
        }
    }
