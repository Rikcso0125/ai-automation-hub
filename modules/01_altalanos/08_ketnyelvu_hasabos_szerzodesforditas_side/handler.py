# -*- coding: utf-8 -*-
import os
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def translate_mock(text: str, target_lang: str) -> str:
    """Intelligens jogi minta-fordító mock módhoz."""
    translations = {
        "A Vállalkozó kötelezettséget vállal a Megrendelő győri telephelyén működő gyártósorok automatizált vezérlőszoftverének korszerűsítésére és a kapcsolódó hardveres elemek szállítására a műszaki leírás szerint.":
            "The Contractor undertakes to modernize the automated control software of the production lines operating at the Client's Győr premises and to deliver the associated hardware components in accordance with the technical specifications.",
        "A felek megállapodnak, hogy a munkadíj összege 45.000 EUR + ÁFA, amely a sikeres átadás-átvételi jegyzőkönyv aláírását követően, 15 napos fizetési határidővel, banki átutalással fizetendő.":
            "The parties agree that the fee shall be EUR 45,000 + VAT, payable by wire transfer within 15 days following the mutual execution of the formal handover protocol.",
        "A felek kötelesek a jelen szerződés teljesítése során tudomásukra jutott minden üzleti, műszaki és pénzügyi információt szigorúan bizalmasan kezelni, és azt harmadik fél részére nem adhatják át.":
            "The parties shall keep strictly confidential all business, technical, and financial information disclosed during the performance of this Agreement and shall not disclose it to any third party.",
        "Egyik fél sem tehető felelőssé olyan kötelezettségszegésért, amelyet az ellenőrzési körén kívül eső elháríthatatlan külső körülmény (vis maior), például háború, sztrájk vagy természeti csapás okozott.":
            "Neither party shall be held liable for any breach of obligation caused by unforeseeable events beyond its reasonable control (force majeure), including war, strikes, or natural disasters."
    }
    return translations.get(text, f"[Translation into {target_lang}]: {text}")

def run(payload: Dict[str, Any], config: Dict[str, Any]) -> Dict[str, Any]:
    title = payload.get("contract_title", "Kétnyelvű Szerződés")
    source_lang = payload.get("source_language", config.get("default_source_lang", "Magyar"))
    target_lang = payload.get("target_language", config.get("default_target_lang", "Angol"))
    paragraphs = payload.get("paragraphs", [])
    gov_rule = config.get("governing_language_rule", "A forrásnyelv (Magyar) az irányadó vita esetén")

    # Fordítás előállítása bekezdésenként
    translated_pairs = []
    for p in paragraphs:
        c_no = p.get("clause_no", "")
        c_title = p.get("title", "")
        src_text = p.get("text", "")
        tgt_text = translate_mock(src_text, target_lang)

        # Záradék cím fordítása
        title_trans = {
            "A Szerződés Tárgya": "Subject Matter of the Agreement",
            "Vállalkozási Díj és Fizetési Feltételek": "Service Fee and Payment Terms",
            "Titoktartási Kötelezettség": "Confidentiality Obligation",
            "Vis Maior": "Force Majeure"
        }.get(c_title, f"{c_title} ({target_lang})")

        translated_pairs.append({
            "clause_no": c_no,
            "source_title": c_title,
            "target_title": title_trans,
            "source_text": src_text,
            "target_text": tgt_text
        })

    # Irányadó nyelvi záradék hozzáadása
    gov_clause_no = f"{len(translated_pairs) + 1}.1"
    gov_src = "A jelen szerződés két nyelven (magyar és angol nyelven) készült. A két változat közötti bármely eltérés vagy értelmezési vita esetén a magyar nyelvű szöveg tekintendő irányadónak."
    gov_tgt = "This Agreement is executed in two languages (Hungarian and English). In case of any discrepancy or dispute of interpretation between the two versions, the Hungarian language version shall prevail."

    translated_pairs.append({
        "clause_no": gov_clause_no,
        "source_title": "Irányadó Nyelv",
        "target_title": "Governing Language",
        "source_text": gov_src,
        "target_text": gov_tgt
    })

    # Word (.docx) fájl generálása 2 hasábos táblázattal
    output_dir = os.path.join(os.path.dirname(__file__), "generated_contracts")
    os.makedirs(output_dir, exist_ok=True)
    docx_filename = "Ketnyelvu_Hasabos_Szerzodes_Export.docx"
    docx_path = os.path.join(output_dir, docx_filename)

    doc = Document()
    doc.add_heading(title, level=1)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = f"{source_lang.upper()} EREDETI"
    hdr_cells[1].text = f"{target_lang.upper()} FORDÍTÁS"

    for p in translated_pairs:
        row_cells = table.add_row().cells
        # Bal hasáb (Forrás)
        p_left = row_cells[0].paragraphs[0]
        p_left.add_run(f"{p['clause_no']}. {p['source_title']}\n").bold = True
        p_left.add_run(p['source_text'])

        # Jobb hasáb (Cél)
        p_right = row_cells[1].paragraphs[0]
        p_right.add_run(f"{p['clause_no']}. {p['target_title']}\n").bold = True
        p_right.add_run(p['target_text'])

    doc.save(docx_path)

    return {
        "status": "success",
        "contract_title": title,
        "languages": {
            "source": source_lang,
            "target": target_lang
        },
        "total_clauses_aligned": len(translated_pairs),
        "governing_language_clause_applied": True,
        "aligned_clauses": translated_pairs,
        "generated_files": {
            "docx_export_path": docx_path,
            "file_name": docx_filename,
            "format": "2-column side-by-side table (.docx)"
        }
    }

async def run_async(payload: Dict[str, Any], config: Dict[str, Any]) -> Dict[str, Any]:
    return run(payload, config)
